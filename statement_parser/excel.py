from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

HEADERS = ["DATE", "MONTH", "SETTLED", "ACCOUNT", "VALUE", "CATEGORY", "TYPE", "DESCRIPTION (optional)"]
WIDTHS = [12, 11, 10, 10, 13, 22, 10, 60]


def to_xlsx(rows) -> bytes:
    """rows: dicts with date, settled, account, value, category, type, description."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, width in enumerate(WIDTHS):
        ws.column_dimensions[chr(ord("A") + i)].width = width
    ws.freeze_panes = "A2"

    for r in rows:
        ws.append([
            r["date"], r["date"].strftime("%B"), bool(r["settled"]), r["account"],
            round(float(r["value"]), 2), r["category"] or None, r["type"] or None,
            r["description"] or None,
        ])
        ws.cell(ws.max_row, 1).number_format = "dd/mm/yyyy"
        ws.cell(ws.max_row, 5).number_format = "#,##0.00;(#,##0.00)"

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
